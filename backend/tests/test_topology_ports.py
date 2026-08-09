"""機器から見たポート割付表。

「この機器は何ポートあって、上はどこで、下に何がいて、どこが空いているか」。
増設のときに最初に聞かれる情報で、結線表（1本1行）からは数えないと出てこない。

一番危ないのは空きポート数を勝手に埋めること。Config を取り込んでいない
機器は総ポート数が分からないので、そこに 0 と書くと「満杯」と読まれる。
"""

import io

from app.models.device import Device, DeviceType
from app.models.topology import Interface, Link
from app.services import topology_ports
from tests.conftest import API


def add_device(db, hostname: str, device_type: DeviceType, ports: list[str] | None = None) -> int:
    device = Device(hostname=hostname, vendor="Cisco", model="C9300-48P", device_type=device_type)
    db.add(device)
    db.flush()
    for name in ports or []:
        db.add(
            Interface(
                device_id=device.id,
                name=name,
                normalized_name=name.lower(),
                enabled=True,
            )
        )
    db.commit()
    return device.id


def connect(db, a: int, a_port: str, b: int, b_port: str) -> None:
    db.add(Link(a_device_id=a, a_interface=a_port, b_device_id=b, b_interface=b_port))
    db.commit()


def test_uplink_and_downlink_are_told_apart(db_session):
    """上位機器へ出ているポートが上り、下にぶら下がる方が下り。"""
    router = add_device(db_session, "rt-01", DeviceType.ROUTER, ["gi0/0/0"])
    switch = add_device(db_session, "sw-01", DeviceType.SWITCH, ["gi1/0/1", "gi1/0/2", "gi1/0/3"])
    ap = add_device(db_session, "ap-01", DeviceType.WIRELESS_AP, ["gi0"])

    connect(db_session, switch, "gi1/0/1", router, "gi0/0/0")
    connect(db_session, switch, "gi1/0/2", ap, "gi0")

    entry = next(d for d in topology_ports.collect(db_session) if d.hostname == "sw-01")
    assert [p.interface for p in entry.uplinks] == ["gi1/0/1"]
    assert [p.interface for p in entry.downlinks] == ["gi1/0/2"]
    assert [p.interface for p in entry.unused] == ["gi1/0/3"]
    assert entry.uplinks[0].peer == "rt-01"


def test_unknown_port_count_is_not_reported_as_zero(db_session):
    """Config 未取り込みの機器で、空き 0 と書かないこと。

    0 と書くと「埋まっている」と読まれ、増設できないと判断されてしまう。
    """
    device = add_device(db_session, "sw-99", DeviceType.SWITCH)
    entry = next(d for d in topology_ports.collect(db_session) if d.device_id == device)
    assert entry.total_ports is None

    data = topology_ports.to_xlsx(db_session)
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(data))["結線図"]
    labels = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value]
    assert any("sw-99" in label and "ポート数不明" in label for label in labels)
    assert not any("sw-99" in label and "空き 0" in label for label in labels)


def test_a_link_on_a_port_the_ledger_does_not_know_is_still_shown(db_session):
    """台帳に無いポートでも、線が引かれていれば表に出すこと。

    図にあって表に無いと、どちらが正しいのか分からなくなる。
    """
    a = add_device(db_session, "sw-a", DeviceType.SWITCH, ["gi1/0/1"])
    b = add_device(db_session, "sw-b", DeviceType.SWITCH, ["gi1/0/1"])
    connect(db_session, a, "gi1/0/48", b, "gi1/0/1")

    entry = next(d for d in topology_ports.collect(db_session) if d.hostname == "sw-a")
    assert "gi1/0/48" in [p.interface for p in entry.ports]


def test_devices_are_ordered_upper_first(db_session):
    """上位から順に並ぶこと。シートをそのまま構成図として読むため。"""
    add_device(db_session, "sw-01", DeviceType.SWITCH)
    add_device(db_session, "rt-01", DeviceType.ROUTER)
    add_device(db_session, "fw-01", DeviceType.FIREWALL)

    assert [d.hostname for d in topology_ports.collect(db_session)] == ["rt-01", "fw-01", "sw-01"]


def test_site_filter(db_session):
    add_device(db_session, "sw-01", DeviceType.SWITCH)
    assert topology_ports.collect(db_session, site_id=999) == []


def test_endpoint_returns_a_workbook(client):
    response = client.get(f"{API}/topology/ports.xlsx")
    assert response.status_code == 200, response.text
    assert response.content[:2] == b"PK"  # xlsx は zip

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["ポート一覧", "結線図"]

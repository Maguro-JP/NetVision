"""Excel の図が、画面や SVG と同じものであること。

以前はセルの罫線で別の図（上位を左、下位を右に並べ直したもの）を描いていた。
同じ台帳から2つの違う図が出ると、どちらを信じればいいのか分からなくなる。
いまは SVG と同じ図を、Excel の図形として置く。

ここで見るのは、図形として入っているか（掴んで動かせるか）、線と箱が
SVG と同じ場所にあるか、そして1枚に収まるようにしてあるか。
"""

import io
import re
import zipfile

from openpyxl import load_workbook

from app.models.device import Device, DeviceType
from app.models.site import Site
from app.models.topology import Link
from app.services import topology_excel, topology_export, topology_view
from app.services.topology import load_diagram
from tests.conftest import API


def build(db) -> int:
    site = Site(name="大阪府営業所", address="大阪府")
    db.add(site)
    db.flush()
    rt = Device(
        hostname="osaka-rt-01",
        vendor="NEC",
        model="IX2215",
        device_type=DeviceType.ROUTER,
        site_id=site.id,
    )
    sw = Device(
        hostname="osaka-sw-01",
        vendor="NEC",
        model="QX-S5528P",
        device_type=DeviceType.SWITCH,
        site_id=site.id,
    )
    db.add_all([rt, sw])
    db.flush()
    db.add(
        Link(
            a_device_id=rt.id,
            a_interface="GigaEthernet0/1",
            b_device_id=sw.id,
            b_interface="GigabitEthernet1/0/25",
        )
    )
    db.commit()
    return site.id


def _drawing(book: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(book)) as archive:
        names = [name for name in archive.namelist() if name.startswith("xl/drawings/")]
        assert names, f"図形の紙が入っていない: {archive.namelist()}"
        return archive.read(names[0]).decode("utf-8")


def test_the_figure_is_made_of_shapes_not_cell_borders(db_session):
    """図形なら、Excel の上で掴んで動かせる。罫線では動かせない。"""
    build(db_session)
    book = topology_excel.to_xlsx(db_session)
    drawing = _drawing(book)
    assert drawing.count("<xdr:sp ") >= 3, "図形が少なすぎる"
    assert "osaka-rt-01" in drawing
    assert "osaka-sw-01" in drawing
    # 結線は折れ線（custGeom）で入る
    assert "custGeom" in drawing


def test_the_port_names_are_in_the_figure(db_session):
    """線のそばのポート名まで含めて、SVG と同じものが入ること。"""
    build(db_session)
    drawing = _drawing(topology_excel.to_xlsx(db_session))
    assert "GE0/1" in drawing or "GigaEthernet0/1" in drawing


def test_the_boxes_are_where_the_svg_puts_them(db_session):
    """同じ台帳から出した2つの図で、箱の場所が違ってはいけない。"""
    build(db_session)
    graph = load_diagram(db_session, "physical")
    svg = topology_export.render_svg(graph, title="構成図")
    boxes = [
        (float(x), float(y))
        for x, y, w in re.findall(r'<rect x="([\d.\-]+)" y="([\d.\-]+)" width="([\d.]+)"', svg)
        if float(w) > 100
    ]
    drawing = _drawing(topology_excel.to_xlsx(db_session))
    spots = {
        (round(int(x) / 9525), round(int(y) / 9525))
        for x, y in re.findall(r'<xdr:pos x="(\d+)" y="(\d+)"/>', drawing)
    }
    for x, y in boxes:
        assert (round(x), round(y)) in spots, f"SVG の箱 {x},{y} が Excel に無い"


def test_the_screen_can_hand_its_own_figure_to_excel(db_session):
    """画面が引いた線は、Excel の図にもそのまま入ること。"""
    build(db_session)
    graph = load_diagram(db_session, "physical")
    edge = graph.edges[0]
    view = topology_view.ViewGeometry(
        wires={
            edge.id: topology_view.ViewWire(
                edge_id=edge.id,
                points=[(0.0, 0.0), (700.0, 0.0), (700.0, 500.0)],
                source_side="right",
                target_side="top",
            )
        }
    )
    drawing = _drawing(topology_excel.to_xlsx(db_session, view=view))
    # 折れ線の幅は、預けた形の幅（700px）と同じになる
    widths = {int(w) for w in re.findall(r'<a:path w="(\d+)"', drawing)}
    assert round(700 * 9525) in widths, f"預けた形の線が無い: {sorted(widths)}"


def test_it_is_set_to_print_on_one_sheet(db_session):
    """紙で配るために出す。既定のままだと A4 の幅で線が切られる。"""
    build(db_session)
    sheet = load_workbook(io.BytesIO(topology_excel.to_xlsx(db_session)))["構成図"]
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
    # 印刷範囲は決めない。決めると Excel が普通の画面に点線の枠を出し、
    # 図の上に大きな枠が乗っているように見える。代わりに、図の右下にあたる
    # セルを作って「ここまでが中身」と伝える
    assert not sheet.print_area
    assert sheet.max_row > 1 and sheet.max_column > 1


def test_tables_are_in_the_same_file(db_session):
    """図と表を別ファイルにしない。配るのは1つで済ませたい。"""
    build(db_session)
    book = load_workbook(io.BytesIO(topology_excel.to_xlsx(db_session)))
    assert book.sheetnames == ["構成図", "結線表", "ポート一覧"]


def test_the_file_still_opens_as_a_workbook(db_session):
    """図形を後から足しているので、ファイルとして壊れていないことを見る。"""
    build(db_session)
    book = load_workbook(io.BytesIO(topology_excel.to_xlsx(db_session)))
    assert book["結線表"]["A1"].value


def test_site_filter_keeps_the_uplink(db_session):
    site_id = build(db_session)
    drawing = _drawing(topology_excel.to_xlsx(db_session, site_id=site_id))
    assert "osaka-rt-01" in drawing


def test_endpoint(client):
    response = client.get(f"{API}/topology/workbook.xlsx")
    assert response.status_code == 200, response.text
    assert response.content[:2] == b"PK"


def test_the_upper_device_is_decided_by_the_whole_network_not_by_link_count(db_session):
    """上下を「繋がっている本数」で決めないこと。

    支社ルーターは配下の拠点が多いぶん本数が多い。本数で並べると、
    センターより支社が上に来てしまう。
    """
    from app.services import topology
    from app.services.topology import hierarchy_rank

    site = Site(name="関西支社", address="大阪府")
    db_session.add(site)
    db_session.flush()

    def router(hostname: str) -> Device:
        device = Device(
            hostname=hostname,
            vendor="Cisco",
            model="ISR4451-X",
            device_type=DeviceType.ROUTER,
            site_id=site.id,
        )
        db_session.add(device)
        return device

    core = router("center-rt-01")
    branch = router("branch-rt-01")
    stations = [router(f"station-{i}-rt-01") for i in range(4)]
    upstream = router("hq-rt-01")
    db_session.flush()

    pairs = [(upstream, core), (core, branch)] + [(branch, s) for s in stations]
    for a, b in pairs:
        db_session.add(
            Link(a_device_id=a.id, a_interface="Gi0/0", b_device_id=b.id, b_interface="Gi0/1")
        )
    db_session.commit()

    rank = hierarchy_rank(topology.get_graph(db_session))
    # 本数は branch(5) > core(2) だが、上位は core 側
    assert rank[core.id] < rank[branch.id]
    assert rank[branch.id] < rank[stations[0].id]

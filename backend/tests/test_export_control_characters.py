"""名前や説明に制御文字が混じっていても、書き出しが止まらない。

端末のログから写した Config の ESC などが1台に混じるだけで、台帳全体の
Excel や構成図が 500 になっていた。書けない文字だけを落とし、残りは出す。
"""

import io
import xml.etree.ElementTree as ET
import zipfile

import pytest
from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from app.services import xlsxsafe
from tests.conftest import API


def test_the_pattern_matches_the_one_openpyxl_refuses() -> None:
    """openpyxl を読み込まずに同じ範囲を持っているので、ずれていないことを見る。"""
    assert xlsxsafe.ILLEGAL_CHARACTERS_RE.pattern == ILLEGAL_CHARACTERS_RE.pattern


def test_only_unwritable_characters_are_dropped() -> None:
    assert xlsxsafe.cell_value("a\x1b[0mb\tc\nd") == "a[0mb\tc\nd"
    assert xlsxsafe.cell_value(42) == 42


@pytest.fixture()
def seeded(client):
    site = client.post(f"{API}/sites", json={"code": "B\x02R", "name": "支\x0b店"}).json()
    response = client.post(
        f"{API}/devices",
        json={
            "hostname": "core\x1b[0m-sw",
            "description": "a\x01b",
            "vendor": "Cisco",
            "model": "C9300-48P",
            "site_id": site["id"],
        },
    )
    assert response.status_code == 201
    return client


@pytest.mark.parametrize(
    "path",
    [
        "/export/devices.xlsx",
        "/export/sites.xlsx",
        "/topology/ports.xlsx",
        "/topology/export.xlsx",
        "/topology/workbook.xlsx",
    ],
)
def test_the_workbooks_are_written(seeded, path) -> None:
    response = seeded.get(f"{API}{path}")
    assert response.status_code == 200
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    for name in archive.namelist():
        if name.endswith(".xml"):
            ET.fromstring(archive.read(name))


def test_the_device_ledger_keeps_the_rest_of_the_name(seeded) -> None:
    response = seeded.get(f"{API}/export/devices.xlsx")
    book = load_workbook(io.BytesIO(response.content))
    values = {cell.value for row in book.active.iter_rows() for cell in row}
    assert "core[0m-sw" in values


def test_the_diagram_is_well_formed_xml(seeded) -> None:
    response = seeded.get(f"{API}/topology/export.svg")
    assert response.status_code == 200
    ET.fromstring(response.content)

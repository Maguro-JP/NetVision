""" "=" で始まる名前や説明を、Excel の式として書き出さない。

openpyxl は "=" で始まる文字列を式に、"#N/A" の形の文字列をエラー値にする。
備考の「=== 要確認 ===」が壊れた式になり、「=HYPERLINK(...)」が開いた人の
Excel でリンクとして動いていた。文字列として書き、読み戻すと同じ値になることを見る。
"""

import io
import zipfile

import pytest
from openpyxl import load_workbook
from openpyxl.cell.cell import ERROR_CODES

from app.services import xlsxsafe
from tests.conftest import API

HYPERLINK = '=HYPERLINK("http://example.invalid","開く")'


def test_the_error_codes_match_the_ones_openpyxl_uses() -> None:
    assert xlsxsafe.ERROR_CODES == frozenset(ERROR_CODES)


def test_plain_text_is_left_as_a_plain_string() -> None:
    assert xlsxsafe.cell_value("core-sw-01") == "core-sw-01"
    assert type(xlsxsafe.cell_value("-")) is str


@pytest.fixture()
def seeded(client):
    site = client.post(f"{API}/sites", json={"code": "=S1", "name": "#N/A"}).json()
    response = client.post(
        f"{API}/devices",
        json={
            "hostname": HYPERLINK,
            "description": "#N/A",
            "vendor": "Cisco",
            "model": "=== 要確認 ===",
            "site_id": site["id"],
        },
    )
    assert response.status_code == 201
    return client


def _sheets(content: bytes) -> list[str]:
    archive = zipfile.ZipFile(io.BytesIO(content))
    return [
        archive.read(name).decode("utf-8")
        for name in archive.namelist()
        if name.startswith("xl/worksheets/sheet")
    ]


@pytest.mark.parametrize(
    "path",
    [
        "/export/devices.xlsx",
        "/export/sites.xlsx",
        "/topology/ports.xlsx",
        "/topology/export.xlsx",
        "/topology/workbook.xlsx",
    ],
)
def test_no_workbook_carries_a_formula_or_an_error_value(seeded, path) -> None:
    response = seeded.get(f"{API}{path}")
    assert response.status_code == 200
    for sheet in _sheets(response.content):
        assert "<f>" not in sheet
        assert 't="e"' not in sheet


def test_the_ledger_reads_back_the_same_text(seeded) -> None:
    response = seeded.get(f"{API}/export/devices.xlsx")
    book = load_workbook(io.BytesIO(response.content))
    values = {cell.value for row in book.active.iter_rows() for cell in row}
    assert {"=== 要確認 ===", HYPERLINK} <= values

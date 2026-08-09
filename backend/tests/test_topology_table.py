"""結線表（Excel / CSV）の書き出し。

図が要らない場面がある。作業指示や検収の突合では、並べ替えと絞り込みのできる
表の方が使いやすい。図と同じ情報が、抜けずに表になっていることを見る。
"""

import io

from tests.conftest import API

CORE = """
hostname core-sw-01
!
Model Number: C9500-40X
!
vlan 10
 name USERS
!
interface GigabitEthernet1/0/1
 description to-acc
 switchport mode trunk
 switchport trunk native vlan 999
 switchport trunk allowed vlan 10,20
!
"""

ACC = """
hostname acc-sw-01
!
Model Number: C9200-48P
!
vlan 10
 name USERS
!
interface GigabitEthernet1/0/48
 description to-core
 switchport mode trunk
 switchport trunk native vlan 1
 switchport trunk allowed vlan 10
!
"""


def setup_pair(client):
    for name, body in (("core-sw-01.txt", CORE), ("acc-sw-01.txt", ACC)):
        client.post(
            f"{API}/configs/upload",
            files=[("files", (name, body.encode(), "text/plain"))],
            data={"create_missing": "true"},
        )
    devices = {d["hostname"]: d["id"] for d in client.get(f"{API}/devices").json()}
    client.post(
        f"{API}/topology/links",
        json={
            "a_device_id": devices["core-sw-01"],
            "a_interface": "GigabitEthernet1/0/1",
            "b_device_id": devices["acc-sw-01"],
            "b_interface": "GigabitEthernet1/0/48",
        },
    )
    return devices


def test_csv_has_one_row_per_link(client):
    setup_pair(client)
    res = client.get(f"{API}/topology/export.csv")

    assert res.status_code == 200
    text = res.content.decode("utf-8-sig")
    lines = [line for line in text.splitlines() if line.strip()]

    assert lines[0].startswith("機器A")
    assert len(lines) == 2  # 見出し + 結線1本
    assert "core-sw-01" in lines[1]
    assert "acc-sw-01" in lines[1]


def test_the_table_carries_the_same_warnings_as_the_diagram(client):
    """図で ⚠ が出るものは、表でも分かること。

    表だけ見て作業する人が見落とすと意味が無い。
    """
    setup_pair(client)
    text = client.get(f"{API}/topology/export.csv").content.decode("utf-8-sig")

    assert "ネイティブVLAN不一致" in text
    # 片側だけ許可している VLAN 20 が挙がる
    # （ネイティブが 999 と 1 で食い違うので、その2つも一緒に出る）
    assert "片側のみ VLAN 1,20,999" in text


def test_shift_jis_is_available(client):
    """Excel が UTF-8 を読めない環境向け。"""
    setup_pair(client)
    res = client.get(f"{API}/topology/export.csv", params={"encoding": "shift_jis"})

    assert res.status_code == 200
    assert "機器A" in res.content.decode("cp932")


def test_xlsx_opens_and_has_a_header(client):
    setup_pair(client)
    res = client.get(f"{API}/topology/export.xlsx")

    assert res.status_code == 200
    assert res.headers["content-disposition"].endswith('.xlsx"')

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(res.content))
    sheet = book["結線表"]
    assert [cell.value for cell in sheet[1]][:3] == ["機器A", "拠点A", "ポートA"]
    assert sheet.max_row == 2
    assert sheet.freeze_panes == "A2"


def test_site_filter_keeps_links_that_only_touch_it_on_one_side(client):
    """拠点をまたぐ回線こそ表で見たいので、片側でも掛かっていれば出す。"""
    devices = setup_pair(client)
    site = client.post(f"{API}/sites", json={"name": "大阪支社"}).json()
    client.patch(f"{API}/devices/{devices['acc-sw-01']}", json={"site_id": site["id"]})

    text = client.get(f"{API}/topology/export.csv", params={"site_id": site["id"]}).content.decode(
        "utf-8-sig"
    )

    assert "acc-sw-01" in text
    assert "core-sw-01" in text  # 相手も行に残る

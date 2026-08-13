"""結線表だけ先に取り込めること。

導入のとき、手元にあるのが結線表だけということがある。機器を1台ずつ
登録し終えるまで結線を入れられないのでは、順番の都合を人に押し付ける
ことになる。書けるものから書いて、揃った時点で繋がるのが正しい。
"""

import io

from openpyxl import load_workbook

from tests.conftest import API

HEADER = "機器A,ポートA,機器B,ポートB,回線名\n"


def upload(client, body: str, *, commit=True, filename="links.csv"):
    path = f"{API}/import/links" if commit else f"{API}/import/links/preview"
    return client.post(path, files={"file": (filename, body.encode("utf-8"), "text/csv")})


def add(client, device_payload, hostname, **kwargs):
    return client.post(f"{API}/devices", json=device_payload(hostname=hostname, **kwargs)).json()[
        "id"
    ]


def test_template_can_be_downloaded_and_filled_in(client):
    """ひな形を出して、そのまま書いて取り込めること。

    列名を当てさせる作りだと、書き始める前に手が止まる。
    """
    res = client.get(f"{API}/import/links/template.xlsx")
    assert res.status_code == 200
    assert res.content[:2] == b"PK"

    book = load_workbook(io.BytesIO(res.content))
    sheet = book["結線表"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == ["機器A", "ポートA", "機器B", "ポートB", "回線名"]
    # 書き方を別紙で持たせる。画面を離れても分かるように
    assert "書き方" in book.sheetnames
    guide = "\n".join(
        str(value or "") for row in book["書き方"].iter_rows(values_only=True) for value in row
    )
    assert "回線名" in guide
    assert "ポート名は機器の脇" in guide


def test_the_example_row_in_the_template_is_not_taken_in(client):
    """ひな形の見本行を消し忘れても、架空の結線にはしない。"""
    res = client.get(f"{API}/import/links/template.xlsx")
    book = load_workbook(io.BytesIO(res.content))
    rows = list(book["結線表"].iter_rows(min_row=2, values_only=True))
    body = HEADER + "\n".join(",".join(str(v or "") for v in row) for row in rows) + "\n"

    result = upload(client, body).json()
    assert result["linked"] == 0
    assert result["pending"] == 0


def test_links_are_created_when_both_devices_exist(client, device_payload):
    add(client, device_payload, "core-sw-01")
    add(client, device_payload, "edge-rt-01", serial_number="X2")

    result = upload(
        client, HEADER + "core-sw-01,GigabitEthernet1/0/1,edge-rt-01,ge-0/0/0,本社 上位\n"
    ).json()

    assert result["linked"] == 1
    assert result["pending"] == 0
    links = client.get(f"{API}/topology/links").json()
    assert len(links) == 1


def test_rows_are_held_until_the_devices_are_registered(client, device_payload):
    """機器より先に結線表を取り込んでも、あとで繋がること。

    ここが繋がらないと、取り込み直しを求めることになる。
    """
    result = upload(client, HEADER + "core-sw-01,Gi1/0/1,edge-rt-01,ge-0/0/0,\n").json()
    assert result["linked"] == 0
    assert result["pending"] == 1
    assert client.get(f"{API}/topology/links").json() == []

    add(client, device_payload, "core-sw-01")
    # 片方だけではまだ繋がらない
    assert client.get(f"{API}/topology/links").json() == []

    add(client, device_payload, "edge-rt-01", serial_number="X2")

    # 取り込み直さなくても結線になっている
    links = client.get(f"{API}/topology/links").json()
    assert len(links) == 1


def test_hostname_case_and_domain_do_not_have_to_match(client, device_payload):
    upload(client, HEADER + "CORE-SW-01.example.jp,Gi1/0/1,edge-rt-01,ge-0/0/0,\n")
    add(client, device_payload, "core-sw-01")
    add(client, device_payload, "edge-rt-01", serial_number="X2")

    assert len(client.get(f"{API}/topology/links").json()) == 1


def test_devices_are_not_invented_from_the_sheet(client):
    """表に書かれた名前から機器を作らない。打ち間違いが台帳に残る方が困る。"""
    upload(client, HEADER + "typo-sw-99,Gi1/0/1,edge-rt-01,ge-0/0/0,\n")
    assert client.get(f"{API}/devices").json() == []


def test_the_same_sheet_twice_does_not_double_up(client, device_payload):
    add(client, device_payload, "core-sw-01")
    add(client, device_payload, "edge-rt-01", serial_number="X2")
    body = HEADER + "core-sw-01,Gi1/0/1,edge-rt-01,ge-0/0/0,\n"

    first = upload(client, body).json()
    second = upload(client, body).json()

    assert first["linked"] == 1
    assert second["linked"] == 0
    assert second["already"] == 1
    assert len(client.get(f"{API}/topology/links").json()) == 1


def test_held_rows_are_not_stored_twice(client, device_payload):
    body = HEADER + "core-sw-01,Gi1/0/1,edge-rt-01,ge-0/0/0,\n"
    upload(client, body)
    upload(client, body)

    add(client, device_payload, "core-sw-01")
    add(client, device_payload, "edge-rt-01", serial_number="X2")
    assert len(client.get(f"{API}/topology/links").json()) == 1


def test_preview_does_not_change_anything(client, device_payload):
    add(client, device_payload, "core-sw-01")
    add(client, device_payload, "edge-rt-01", serial_number="X2")

    result = upload(
        client, HEADER + "core-sw-01,Gi1/0/1,edge-rt-01,ge-0/0/0,\n", commit=False
    ).json()

    assert result["linked"] == 1  # 取り込めば1本になる、という見立て
    assert result["committed"] is False
    assert client.get(f"{API}/topology/links").json() == []


def test_incomplete_rows_are_reported_with_their_line_number(client):
    body = HEADER + "core-sw-01,Gi1/0/1,,,\n" + "a-sw,Gi1,a-sw,Gi2,自分どうし\n"
    result = upload(client, body).json()

    assert [issue["row_number"] for issue in result["issues"]] == [2, 3]
    assert "4つとも" in result["issues"][0]["message"]
    assert "同じ機器" in result["issues"][1]["message"]


def test_a_sheet_without_the_needed_columns_is_rejected_clearly(client):
    result = upload(client, "適当,列\n1,2\n").json()
    assert result["issues"]
    assert "機器A" in result["issues"][0]["message"]


def test_columns_from_a_real_world_sheet_are_understood(client, device_payload):
    """現場の表をそのまま持ってきても読めること。"""
    add(client, device_payload, "core-sw-01")
    add(client, device_payload, "edge-rt-01", serial_number="X2")

    body = "自機器,自ポート,対向機器,対向ポート,回線\n"
    body += "core-sw-01,Gi1/0/1,edge-rt-01,ge-0/0/0,本社\n"

    assert upload(client, body).json()["linked"] == 1


def test_devices_added_by_the_inventory_import_also_resolve_links(client):
    """台帳の取り込みで機器が増えた場合も繋がること。"""
    upload(client, HEADER + "imp-sw-01,Gi1/0/1,imp-rt-01,ge-0/0/0,\n")

    csv = "ホスト名,メーカー,型番\nimp-sw-01,Cisco,C9300-48P\nimp-rt-01,Juniper,MX204\n"
    client.post(
        f"{API}/import/devices",
        files={"file": ("devices.csv", csv.encode("utf-8"), "text/csv")},
    )

    assert len(client.get(f"{API}/topology/links").json()) == 1

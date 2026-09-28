"""CSV / Excel ファイルを「ヘッダ + 行の辞書」に読み込むユーティリティ。

既存の Excel 台帳をそのまま取り込めるようにするのが目的なので、
日本語環境で実際に出てくる形式（CP932 の CSV、BOM 付き UTF-8、.xlsx）を許容する。
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime

from app.services import csvsafe, textencoding

#: 判別は textencoding にまとめてある。cp932 と EUC-JP はお互いのバイト列を
#: エラー無しに読めるので、順に試して最初に通ったものを採ると化ける
CSV_ENCODINGS = textencoding.ENCODINGS


class TableFileError(ValueError):
    """ファイルとして読めなかった場合に送出する。"""


@dataclass
class Table:
    headers: list[str]
    rows: list[dict[str, object]] = field(default_factory=list)
    #: 最初のデータ行が、ファイルの何行目か。見出しの上に表題の行があると 2 ではなくなる。
    #: 取り込みの結果で「何行目」と言うときの起点にする（人が Excel で開いて探せるように）
    first_row: int = 2


ZIP_MAGIC = b"PK\x03\x04"


def read_table(filename: str, content: bytes) -> Table:
    name = (filename or "").lower()
    # .xlsx は ZIP なので、拡張子が付いていない/変えられていても中身で判別できる
    if content.startswith(ZIP_MAGIC) or name.endswith((".xlsx", ".xlsm")):
        return _read_excel(content)
    if name.endswith((".csv", ".txt", ".tsv")):
        return _read_csv(content, delimiter="\t" if name.endswith(".tsv") else None)
    raise TableFileError("対応していない拡張子です（.csv / .tsv / .xlsx に対応しています）")


def _decode(content: bytes) -> str:
    # 表は、化けた文字が1つ入るだけで台帳として使えなくなる。読めた分だけ
    # 通すより、読めなかったと言って保存し直してもらうほうがよい
    text = textencoding.best_effort(content)
    if text is None:
        raise TableFileError(
            "文字コードを判別できませんでした（UTF-8 または Shift_JIS で保存してください）"
        )
    return text


def _read_csv(content: bytes, delimiter: str | None = None) -> Table:
    text = _decode(content)
    if not text.strip():
        raise TableFileError("ファイルが空です")

    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",\t;").delimiter
        except csv.Error:
            delimiter = ","

    lines, records = _csv_records(text, delimiter)
    if not records:
        raise TableFileError("データ行がありません")
    header_at = _header_index(records)
    first_row = lines[header_at] + 1
    records = records[header_at:]

    headers = [cell.strip() for cell in records[0]]
    rows: list[dict[str, object]] = []
    for record in records[1:]:
        row: dict[str, object] = {}
        for index, header in enumerate(headers):
            # 書き出しで式よけに付けた先頭の `'` を外す（csvsafe.py）
            row[header] = csvsafe.unguard(record[index].strip()) if index < len(record) else ""
        rows.append(row)
    return Table(headers=headers, rows=rows, first_row=first_row)


def _csv_records(text: str, delimiter: str) -> tuple[list[int], list[list[str]]]:
    """CSV の行を読む。閉じていない「"」を黙って通さない。

    欄の中に閉じていない「"」が1つあると、そこから後ろが全部1つの欄になる。
    ゆるい読み方ではそのまま通り、数千行の台帳が「5行・エラー0件」として取り込まれ、
    残りが黙って消える。厳密に読んで「引用の途中でファイルが終わった」ことを見つけ、
    どの行かを返す。それ以外の厳密さの違い（`"x"y` のような書き方）は、今までどおり
    ゆるく読む（手で作った CSV を急に弾かないため）。
    """
    lines: list[int] = []
    records: list[list[str]] = []
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    try:
        for row in reader:
            if any(cell.strip() for cell in row):
                lines.append(reader.line_num)
                records.append(row)
        return lines, records
    except csv.Error as exc:
        started = reader.line_num
        if "unexpected end of data" in str(exc):
            raise TableFileError(
                f"CSV を最後まで読めませんでした。見出しから数えて{len(records) + 1}件目の行で、"
                f'欄の中の「"」が閉じていません（ファイルの終わりの{started}行目まで続いています）。'
                '閉じていない「"」を直してから取り込んでください'
            ) from exc
        if "field larger than field limit" in str(exc):
            raise TableFileError(
                f"見出しから数えて{len(records) + 1}件目の行に、"
                "大きすぎる欄（128KB 超）があります。"
                '閉じていない「"」で後ろの行まで1つの欄になっていないか確かめてください'
            ) from exc

    loose = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    lines, records = [], []
    try:
        for row in loose:
            if any(cell.strip() for cell in row):
                lines.append(loose.line_num)
                records.append(row)
        return lines, records
    except csv.Error as exc:
        raise TableFileError(f"CSV を読めませんでした（{loose.line_num}行目付近: {exc}）") from exc


def _read_excel(content: bytes) -> Table:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - 依存が入っていない環境向け
        raise TableFileError("Excel の読み込みには openpyxl が必要です") from exc

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise TableFileError(f"Excel ファイルを読めませんでした: {exc}") from exc

    sheet = workbook.active
    if sheet is None:
        raise TableFileError("シートが見つかりません")

    numbered = [
        (number, row)
        for number, row in enumerate(sheet.iter_rows(values_only=True), start=1)
        if any(_clean(c) for c in row)
    ]
    workbook.close()
    if not numbered:
        raise TableFileError("データ行がありません")
    records = [row for _number, row in numbered]
    header_at = _header_index(records)
    first_row = numbered[header_at][0] + 1
    records = records[header_at:]

    headers = [_clean(cell) for cell in records[0]]
    rows: list[dict[str, object]] = []
    for record in records[1:]:
        row: dict[str, object] = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            value = record[index] if index < len(record) else None
            # Excel の日付セルは datetime で返るので、日付として扱えるようそのまま渡す
            row[header] = value if isinstance(value, (datetime, date)) else _clean(value)
        rows.append(row)
    return Table(headers=[h for h in headers if h], rows=rows, first_row=first_row)


#: 見出しを探す範囲。表題や作成日の行が、これより多く続く台帳は見ない
_HEADER_SEARCH_ROWS = 10


def _header_index(records: list) -> int:
    """見出しの行が何行目か（0 始まり）。

    日本の台帳は、1行目に表題（「機器台帳（2026年度）」など）を置き、その下に
    見出しを書くことが多い。先頭の行をそのまま見出しにすると、どの列も対応付かず
    「ホスト名の列がありません」になり、表題の行を手で消すまで取り込めない。
    埋まっている欄が1つ以下の行（表題、空行）は、後ろに2欄以上の行があれば
    見出しとみなさない。どの語が見出しかは当て推量しない（取り込みごとに語が違うため）。
    """
    for index, record in enumerate(records[:_HEADER_SEARCH_ROWS]):
        if sum(1 for cell in record if _clean(cell)) >= 2:
            return index
    return 0


def _clean(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()

"""機器ごとのポート割付を、結線図として書き出す。

結線表（1本1行）は「線」から見た表で、工事の指示には向く。一方で現場で
欲しくなるのは**機器から見た表**である。この機器は何ポートあって、上位の
どこへ上がっていて、下に何がぶら下がっていて、どこが空いているのか。
増設のときに真っ先に聞かれるのはこれで、線の一覧からは数えないと分からない。

並びは上位を左・下位を右にする。機器の左にアップリンク、右にダウンリンクを
置くと、シートをそのまま構成図として読める。

**空きポートは、台帳に載っているポートからしか出せない。** Config を
取り込んでいない機器は総ポート数が分からないので、そこは空欄にして
「0ポート」とは書かない。数字が入っていると数えたように読まれてしまう。
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.device import Device
from app.models.topology import Interface, Link
from app.services import portlabel
from app.services.configparse import ifname
from app.services.topology import TIER_ORDER

UPLINK = "上り"
DOWNLINK = "下り"
LATERAL = "横"
UNUSED = "空き"

HEADERS = (
    "機器",
    "拠点",
    "機種",
    "ポート",
    "種別",
    "区分",
    "接続先",
    "接続先ポート",
    "状態",
    "VLAN",
    "VRF",
    "説明",
)


@dataclass
class PortRow:
    interface: str
    media: str
    role: str
    peer: str = ""
    peer_port: str = ""
    enabled: bool = True
    vlan: str = ""
    vrf: str = ""
    description: str = ""


@dataclass
class DevicePorts:
    device_id: int
    hostname: str
    site_name: str
    model: str
    tier: int
    # 台帳に載っているポート数。Config 未取り込みなら None（0 とは書かない）
    total_ports: int | None
    ports: list[PortRow] = field(default_factory=list)

    @property
    def uplinks(self) -> list[PortRow]:
        return [p for p in self.ports if p.role == UPLINK]

    @property
    def downlinks(self) -> list[PortRow]:
        return [p for p in self.ports if p.role in (DOWNLINK, LATERAL)]

    @property
    def unused(self) -> list[PortRow]:
        return [p for p in self.ports if p.role == UNUSED]


def _sort_key(name: str) -> tuple:
    """GigabitEthernet1/0/2 が 1/0/10 より前に来るように、数字は数として見る。"""
    _, number, _ = portlabel.split(name)
    parts = [int(p) for p in (number or "").replace(":", "/").split("/") if p.isdigit()]
    return (portlabel.media_of(name).value, parts, name)


def collect(db: Session, *, site_id: int | None = None) -> list[DevicePorts]:
    devices = list(db.scalars(select(Device)))
    by_id = {d.id: d for d in devices}

    interfaces: dict[int, list[Interface]] = {}
    for interface in db.scalars(select(Interface)):
        interfaces.setdefault(interface.device_id, []).append(interface)

    # 結線を「機器 → 自分のポート → 相手」に組み替える
    linked: dict[int, dict[str, tuple[int, str]]] = {}
    for link in db.scalars(select(Link)):
        linked.setdefault(link.a_device_id, {})[ifname.normalize(link.a_interface)] = (
            link.b_device_id,
            link.b_interface,
        )
        linked.setdefault(link.b_device_id, {})[ifname.normalize(link.b_interface)] = (
            link.a_device_id,
            link.a_interface,
        )

    out: list[DevicePorts] = []
    for device in devices:
        if site_id is not None and device.site_id != site_id:
            continue

        tier = TIER_ORDER.get(device.device_type, 7)
        mine = linked.get(device.id, {})
        known = interfaces.get(device.id, [])
        entry = DevicePorts(
            device_id=device.id,
            hostname=device.hostname,
            site_name=device.site.name if device.site else "",
            model=device.model,
            tier=tier,
            total_ports=len(known) or None,
        )

        seen: set[str] = set()
        for interface in known:
            # 突き合わせは必ず同じ関数で正規化する。台帳の normalized_name は
            # 取り込み時の形なので、線の側と綴りが揃っているとは限らない
            key = ifname.normalize(interface.name)
            seen.add(key)
            peer = mine.get(key)
            entry.ports.append(_row(interface.name, peer, by_id, tier, interface=interface))

        # 台帳に無いポートでも、線が引かれているなら出す。
        # 図に描いてあるものが表に無いと、どちらが正しいのか分からなくなる
        for key, peer in mine.items():
            if key in seen:
                continue
            entry.ports.append(_row(peer_name(db, device.id, key) or key, peer, by_id, tier))

        entry.ports.sort(key=lambda row: _sort_key(row.interface))
        out.append(entry)

    # 上位を先に。同じ段は名前順
    out.sort(key=lambda d: (d.tier, d.site_name, d.hostname))
    return out


def peer_name(db: Session, device_id: int, normalized: str) -> str | None:
    """線に書かれているポート名をそのまま拾う（台帳に無いポート用）。"""
    for link in db.scalars(select(Link).where(Link.a_device_id == device_id)):
        if ifname.normalize(link.a_interface) == normalized:
            return link.a_interface
    for link in db.scalars(select(Link).where(Link.b_device_id == device_id)):
        if ifname.normalize(link.b_interface) == normalized:
            return link.b_interface
    return None


def _row(
    name: str,
    peer: tuple[int, str] | None,
    by_id: dict[int, Device],
    tier: int,
    interface: Interface | None = None,
) -> PortRow:
    row = PortRow(
        interface=name,
        media=portlabel.MEDIA_LABEL.get(portlabel.media_of(name), ""),
        role=UNUSED,
    )
    if interface is not None:
        row.enabled = interface.enabled
        row.vrf = interface.vrf or ""
        row.description = interface.description or ""
        row.vlan = interface.trunk_vlans or (str(interface.vlan) if interface.vlan else "")

    if peer is None:
        return row

    other = by_id.get(peer[0])
    row.peer = other.hostname if other else ""
    row.peer_port = peer[1]
    if other is None:
        row.role = LATERAL
        return row

    other_tier = TIER_ORDER.get(other.device_type, 7)
    # 段が上の相手へ出ているものが上り。同じ段は横（冗長やスタック間）
    row.role = UPLINK if other_tier < tier else DOWNLINK if other_tier > tier else LATERAL
    return row


def rows(db: Session, *, site_id: int | None = None) -> list[list[str]]:
    """1ポート1行の一覧。並べ替え・絞り込みができる形。"""
    out: list[list[str]] = []
    for device in collect(db, site_id=site_id):
        for port in device.ports:
            out.append(
                [
                    device.hostname,
                    device.site_name,
                    device.model,
                    port.interface,
                    port.media,
                    port.role,
                    port.peer,
                    port.peer_port,
                    "有効" if port.enabled else "shutdown",
                    port.vlan,
                    port.vrf,
                    port.description,
                ]
            )
    return out


def to_xlsx(db: Session, *, site_id: int | None = None) -> bytes:
    """「ポート一覧」と「結線図」の2枚を書き出す。

    一覧は絞り込み用、結線図は貼って見せる用。同じ元データから作る。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    devices = collect(db, site_id=site_id)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ポート一覧"
    sheet.append(list(HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    data = rows(db, site_id=site_id)
    for row in data:
        sheet.append(row)

    # 空きポートに色を付ける。増設のときに真っ先に見る列
    free = PatternFill("solid", start_color="E8F4E8", end_color="E8F4E8")
    role_column = HEADERS.index("区分") + 1
    for index in range(2, len(data) + 2):
        if sheet.cell(row=index, column=role_column).value == UNUSED:
            for column in range(1, len(HEADERS) + 1):
                sheet.cell(row=index, column=column).fill = free

    for index, header in enumerate(HEADERS, start=1):
        width = max(len(header) * 2 + 2, 12)
        for row in data:
            width = max(width, min(len(str(row[index - 1])) + 2, 40))
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    # --- 結線図 ---
    tree = workbook.create_sheet("結線図")
    _write_tree(tree, devices, Alignment, Border, Font, PatternFill, Side, get_column_letter)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


_TREE_HEADERS = ("上りポート", "上位機器", "機器", "下りポート", "接続先")


def _write_tree(
    sheet, devices: list[DevicePorts], Alignment, Border, Font, PatternFill, Side, get_column_letter
) -> None:
    """機器を段の順に縦に並べ、左に上り・右に下りを書く。

    左が上位、右が下位。1機器が数行になるので、機器名の列は縦結合する。
    """
    sheet.append(list(_TREE_HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    tier_fill = PatternFill("solid", start_color="EEF3FA", end_color="EEF3FA")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_index = 2
    for device in devices:
        ups, downs = device.uplinks, device.downlinks
        height = max(len(ups), len(downs), 1)

        for offset in range(height):
            up = ups[offset] if offset < len(ups) else None
            down = downs[offset] if offset < len(downs) else None
            sheet.cell(row=row_index + offset, column=1, value=up.interface if up else "")
            sheet.cell(row=row_index + offset, column=2, value=up.peer if up else "")
            sheet.cell(row=row_index + offset, column=4, value=down.interface if down else "")
            sheet.cell(row=row_index + offset, column=5, value=down.peer if down else "")
            for column in range(1, 6):
                sheet.cell(row=row_index + offset, column=column).border = box

        # 総ポート数が分からない機器は、空き数を書かずに「不明」と出す。
        # 0 と書くと「埋まっている」と読まれる
        if device.total_ports is None:
            summary = "ポート数不明（Config 未取り込み）"
        else:
            summary = f"{device.total_ports}ポート / 空き {len(device.unused)}"
        label = f"{device.hostname}\n{device.model}\n{summary}"
        cell = sheet.cell(row=row_index, column=3, value=label)
        cell.alignment = center
        cell.fill = tier_fill
        cell.border = box
        if height > 1:
            sheet.merge_cells(
                start_row=row_index, start_column=3, end_row=row_index + height - 1, end_column=3
            )
        row_index += height

    for column, width in zip((1, 2, 3, 4, 5), (18, 20, 30, 18, 20), strict=True):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"

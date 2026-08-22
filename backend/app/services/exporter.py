"""機器台帳の書き出し。

列見出しは取り込み側（importer.FIELD_LABELS）と揃えてあるので、
書き出したファイルをそのまま取り込み直せる（往復できる）。
Excel で編集して戻す運用や、稟議・見積依頼用の資料作成を想定している。
"""

from __future__ import annotations

import csv
import io
from datetime import date

from app.models import Device, MaintenanceContract, Site
from app.services.importer import FIELD_LABELS, IMPORTABLE_FIELDS
from app.services.lifecycle import LifecycleStatus, Severity, evaluate

# 取り込みでは使わないが、資料として役に立つ算出列
STATUS_COLUMNS: tuple[tuple[str, str], ...] = (
    ("maintenance_end", "保守期限(算出)"),
    ("days_to_maintenance_end", "保守期限まで(日)"),
    ("days_to_eol", "EoLまで(日)"),
    ("age_years", "経過年数"),
    ("overall_severity", "アラート"),
)

SEVERITY_LABELS: dict[Severity, str] = {
    Severity.EXPIRED: "期限切れ",
    Severity.CRITICAL: "要対応",
    Severity.WARNING: "要計画",
    Severity.OK: "問題なし",
    Severity.UNKNOWN: "情報未登録",
}

# 取り込み時に受け付ける日本語表記のうち、書き出しで使う代表表記を逆引きする
_TYPE_LABELS = {
    "switch": "スイッチ",
    "router": "ルーター",
    "firewall": "ファイアウォール",
    "load_balancer": "ロードバランサー",
    "wireless_ap": "無線AP",
    "wireless_controller": "無線コントローラー",
    "server": "サーバー",
    "other": "その他",
}
_STATUS_LABELS = {
    "active": "稼働中",
    "standby": "予備機",
    "maintenance": "保守中",
    "retired": "撤去済み",
}

# これらの表記が取り込み側でも解釈できることは tests/test_export_api.py で担保している


def build_headers(include_status: bool = True) -> list[str]:
    headers = [FIELD_LABELS[f] for f in IMPORTABLE_FIELDS]
    if include_status:
        headers += [label for _, label in STATUS_COLUMNS]
    return headers


def _latest_contract(device: Device) -> MaintenanceContract | None:
    contracts = [c for c in device.contracts if c.end_date is not None]
    return max(contracts, key=lambda c: c.end_date) if contracts else None


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, date):
        return value.strftime("%Y/%m/%d")
    return str(value)


def build_row(device: Device, include_status: bool = True) -> list[object]:
    """1台分の行を作る。保守契約は最も終了日が遠いものを代表として書き出す。"""
    contract = _latest_contract(device)
    status: LifecycleStatus = evaluate(device)

    values: dict[str, object] = {
        "device_id": device.id,
        "code": device.code,
        "label": device.label,
        "hostname": device.hostname,
        "vendor": device.vendor,
        "model": device.model,
        "serial_number": device.serial_number,
        "device_type": _TYPE_LABELS.get(str(device.device_type), ""),
        "status": _STATUS_LABELS.get(str(device.status), ""),
        "os_name": device.os_name,
        "os_version": device.os_version,
        "firmware_version": device.firmware_version,
        "site": device.site.name if device.site else None,
        "site_code": device.site.code if device.site else None,
        "rack": device.rack,
        "rack_unit": device.rack_unit,
        "mgmt_ip": device.mgmt_ip,
        "install_date": device.install_date,
        "eos_date": device.eos_date,
        "eol_date": device.eol_date,
        "notes": device.notes,
        "contract_provider": contract.provider if contract else None,
        "contract_number": contract.contract_number if contract else None,
        "contract_service_level": contract.service_level if contract else None,
        "contract_start_date": contract.start_date if contract else None,
        "contract_end_date": contract.end_date if contract else None,
        "contract_annual_cost": (
            float(contract.annual_cost) if contract and contract.annual_cost is not None else None
        ),
    }

    row: list[object] = [values[f] for f in IMPORTABLE_FIELDS]
    if include_status:
        row += [
            status.maintenance_end,
            status.days_to_maintenance_end,
            status.days_to_eol,
            status.age_years,
            SEVERITY_LABELS[status.overall_severity],
        ]
    return row


def to_csv(
    devices: list[Device], *, include_status: bool = True, encoding: str = "utf-8-sig"
) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(build_headers(include_status))
    for device in devices:
        writer.writerow([_text(v) for v in build_row(device, include_status)])
    # Shift_JIS に無い文字（絵文字や一部の記号）で落ちないよう、変換できない文字は代替する
    return buffer.getvalue().encode(encoding, errors="replace")


def stream_csv(
    devices, *, include_status: bool = True, encoding: str = "utf-8-sig", chunk: int = 2000
):
    """CSV を少しずつ返す。

    16万台では、全部を組み立ててから返すと最初の1バイトが出るまで40秒かかる。
    その間ブラウザは何も表示しないので、押しても反応が無いように見える。
    数千行ごとに返せば、押した直後から保存が始まる。
    """
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    # 印（BOM）は先頭に1つだけ。塊ごとに付けると、その都度ゴミの行が増える
    body_encoding = "utf-8" if encoding == "utf-8-sig" else encoding

    def take() -> bytes:
        text = buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)
        # Shift_JIS に無い文字（絵文字や一部の記号）で落ちないよう、代替する
        return text.encode(body_encoding, errors="replace")

    if encoding == "utf-8-sig":
        yield "\ufeff".encode()

    writer.writerow(build_headers(include_status))
    yield take()

    for index, device in enumerate(devices, start=1):
        writer.writerow([_text(v) for v in build_row(device, include_status)])
        if index % chunk == 0:
            yield take()
    rest = take()
    if rest:
        yield rest


def to_xlsx(devices, *, include_status: bool = True) -> bytes:
    """機器台帳を Excel で書き出す。

    16万台では、ふつうの組み立て方だと全行をメモリに持ったまま最後に書くので、
    1分以上かかる。openpyxl の「書き出し専用」で、行を書いたそばから流す。
    そのぶん、書いたあとにセルへ戻って飾りを足すことはできないので、
    列幅と日付の書式は最初に決めておく。
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("機器台帳")
    # 見出しの固定は、行を書く前に決める（書き出し専用ではあとから戻れない）
    sheet.freeze_panes = "A2"

    headers = build_headers(include_status)
    date_headers = {
        FIELD_LABELS[f]
        for f in (
            "install_date",
            "eos_date",
            "eol_date",
            "contract_start_date",
            "contract_end_date",
        )
    } | {STATUS_COLUMNS[0][1]}

    # 列幅と日付の書式は、行を書く前に決める（あとから戻れないため）
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter] = ColumnDimension(
            sheet, index=letter, width=min(max(len(header) * 2, 12), 40)
        )

    head = []
    for header in headers:
        cell = WriteOnlyCell(sheet, value=header)
        cell.font = Font(bold=True)
        head.append(cell)
    sheet.append(head)

    date_columns = {index for index, header in enumerate(headers) if header in date_headers}
    for device in devices:
        row = []
        for index, value in enumerate(build_row(device, include_status)):
            if index in date_columns and value is not None:
                cell = WriteOnlyCell(sheet, value=value)
                cell.number_format = "yyyy/mm/dd"
                row.append(cell)
            else:
                row.append(value)
        sheet.append(row)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- 拠点の台帳 ------------------------------------------------------------
#
# 機器と同じく、書き出したものを直して戻せる並びにする。拠点IDを先に置くのは、
# 取り込みで拠点を突き合わせる鍵がこの列だから。

SITE_COLUMNS: tuple[tuple[str, str], ...] = (
    ("code", "拠点ID"),
    ("name", "拠点名"),
    ("status", "状態"),
    ("address", "住所"),
    ("opened_on", "開設日"),
    ("closed_on", "閉鎖日"),
    ("device_count", "機器数"),
    ("notes", "メモ"),
)

SITE_STATUS_LABELS = {
    "active": "稼働中",
    "planned": "開設予定",
    "suspended": "一時閉鎖",
    "closed": "閉鎖・閉店",
}


def site_headers() -> list[str]:
    return [label for _, label in SITE_COLUMNS]


def build_site_row(site: Site, device_count: int) -> list[object]:
    values: dict[str, object] = {
        "code": site.code,
        "name": site.name,
        "status": SITE_STATUS_LABELS.get(str(site.status), str(site.status)),
        "address": site.address,
        "opened_on": site.opened_on,
        "closed_on": site.closed_on,
        "device_count": device_count,
        "notes": site.notes,
    }
    return [values[name] for name, _ in SITE_COLUMNS]


def sites_to_csv(rows: list[tuple[Site, int]], *, encoding: str = "utf-8-sig") -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(site_headers())
    for site, count in rows:
        writer.writerow([_text(value) for value in build_site_row(site, count)])
    return buffer.getvalue().encode(encoding, errors="replace")


def sites_to_xlsx(rows: list[tuple[Site, int]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "拠点台帳"
    headers = site_headers()
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for site, count in rows:
        sheet.append(build_site_row(site, count))

    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = min(max(len(header) * 3, 12), 40)
        if header in {"開設日", "閉鎖日"}:
            for cell in sheet[letter][1:]:
                cell.number_format = "yyyy/mm/dd"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

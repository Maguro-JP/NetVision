"""トポロジを表として書き出す。

図が要らない場面がある。**結線表**（どの機器のどのポートが、どこに繋がっているか）は
現場での作業指示や、工事業者への依頼、検収の突合に使う。そこで欲しいのは絵ではなく
並べ替え・絞り込みのできる表で、Excel の方が向いている。

図（SVG）と表（Excel / CSV）は用途が違うので、両方出せるようにしてある。
"""

from __future__ import annotations

import csv
import io

from sqlalchemy.orm import Session

from app.models.topology import LinkOrigin
from app.services import portlabel
from app.services.configparse import vlanlist
from app.services.topology import DEFAULT_LAYER, get_graph

HEADERS = (
    "機器A",
    "拠点A",
    "ポートA",
    "種別",
    "機器B",
    "拠点B",
    "ポートB",
    "確認",
    "取得元",
    "VLAN",
    "ネイティブVLAN",
    "注意",
    "備考",
)

_ORIGIN_LABEL = {LinkOrigin.DISCOVERED: "自動推定", LinkOrigin.MANUAL: "手動"}


def rows(db: Session, layer: str = DEFAULT_LAYER, *, site_id: int | None = None) -> list[list[str]]:
    """結線を1本1行にする。図と同じ情報を、並べ替えられる形で出す。"""
    graph = get_graph(db, layer)
    nodes = {n.device_id: n for n in graph.nodes}

    out: list[list[str]] = []
    for edge in graph.edges:
        a, b = nodes.get(edge.source), nodes.get(edge.target)
        if a is None or b is None:
            continue
        # 拠点で絞るときは、その拠点に**片側でも**掛かっていれば出す。
        # 拠点をまたぐ回線こそ表で確認したいため
        if site_id is not None and site_id not in (a.site_id, b.site_id):
            continue

        warnings: list[str] = []
        if edge.native_mismatch:
            warnings.append("ネイティブVLAN不一致")
        if edge.vlan_mismatch:
            warnings.append(f"片側のみ VLAN {vlanlist.format_ranges(edge.vlan_mismatch)}")

        out.append(
            [
                a.hostname,
                a.site_name or "",
                edge.source_interface,
                portlabel.MEDIA_LABEL.get(portlabel.media_of(edge.source_interface), ""),
                b.hostname,
                b.site_name or "",
                edge.target_interface,
                "両側" if edge.confirmed_both_ways else "片側のみ",
                _ORIGIN_LABEL.get(edge.origin, str(edge.origin)),
                vlanlist.format_ranges(edge.vlans),
                str(edge.native_vlan) if edge.native_vlan else "",
                " / ".join(warnings),
                edge.label or "",
            ]
        )

    out.sort(key=lambda row: (row[1], row[0], row[2]))
    return out


def to_csv(
    db: Session, layer: str = DEFAULT_LAYER, *, site_id: int | None = None, encoding: str = "utf-8"
) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(HEADERS)
    writer.writerows(rows(db, layer, site_id=site_id))

    text = buffer.getvalue()
    if encoding == "shift_jis":
        # 変換できない文字で落とすより、置き換えて出す方がまし
        return text.encode("cp932", errors="replace")
    # Excel が UTF-8 と分かるように BOM を付ける
    return text.encode("utf-8-sig")


def to_xlsx(db: Session, layer: str = DEFAULT_LAYER, *, site_id: int | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "結線表"

    sheet.append(list(HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    data = rows(db, layer, site_id=site_id)
    for row in data:
        sheet.append(row)

    # 食い違いのある行は色を付ける。表でも見落とさないようにする
    warning_column = HEADERS.index("注意") + 1
    highlight = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
    for index in range(2, len(data) + 2):
        if sheet.cell(row=index, column=warning_column).value:
            for column in range(1, len(HEADERS) + 1):
                sheet.cell(row=index, column=column).fill = highlight

    for index, header in enumerate(HEADERS, start=1):
        width = max(len(header) * 2 + 2, 12)
        for row in data:
            width = max(width, min(len(str(row[index - 1])) + 2, 40))
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
